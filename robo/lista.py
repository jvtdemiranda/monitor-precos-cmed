"""
Leitura e limpeza da lista oficial de preços da CMED (arquivo "PMC - xls").

O arquivo publicado pela Anvisa não é uma tabela pronta:
- ~40 linhas de notas antes do cabeçalho (a posição pode mudar);
- todos os valores vêm como texto, inclusive os preços ("1234,56");
- preço com asterisco ("54,23*") marca produto isento de ICMS;
- campo vazio vem como "    -     " ou "- (*) ";
- produto de uso restrito a hospital vem sem PMC em todas as alíquotas;
- códigos de barras repetidos entre produtos, e "0000000000000" como
  código de barras de mentira.

ler_lista() devolve os produtos já limpos e um registro do que foi
corrigido — que aparece na página, na seção "Qualidade da lista oficial".
"""
import re
import unicodedata
from dataclasses import dataclass, field

from openpyxl import load_workbook


class ListaInvalida(Exception):
    """O arquivo não tem o formato esperado — melhor parar do que publicar errado."""


def normalizar(texto):
    """'CÓDIGO  GGREM\xa0' -> 'CODIGO GGREM' (sem acento, maiúsculo, espaços simples)."""
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return " ".join(t.upper().split())


# coluna do arquivo (normalizada) -> campo nosso
OBRIGATORIAS = {
    "CODIGO GGREM": "ggrem",
    "SUBSTANCIA": "substancia",
    "LABORATORIO": "laboratorio",
    "PRODUTO": "produto",
    "APRESENTACAO": "apresentacao",
    "EAN 1": "ean1",
    "RESTRICAO HOSPITALAR": "hospitalar",
}
OPCIONAIS = {
    "CNPJ": "cnpj",
    "REGISTRO": "registro",
    "EAN 2": "ean2",
    "EAN 3": "ean3",
    "CLASSE TERAPEUTICA": "classe",
    "TIPO DE PRODUTO (STATUS DO PRODUTO)": "tipo",
    "REGIME DE PRECO": "regime",
    "ICMS 0%": "icms0",
    "ANALISE RECURSAL": "recursal",
    "TARJA": "tarja",
}
# "PMC 19 %", "PMC 17,5 %  ALC", "PMC 0 %", "PMC Sem Impostos"
RE_PMC = re.compile(r"^PMC (?:(\d+(?:,\d+)?) ?%|(SEM IMPOSTOS))( ALC)?$")
RE_PRECO = re.compile(r"^(\d+),(\d{2})(\*?)$")
RE_VAZIO = re.compile(r"^-?\s*(\(\*\))?$")
RE_PUBLICADA = re.compile(r"Publicada em (\d{2})/(\d{2})/(\d{4})")

# a partir de quantas linhas o arquivo passa a parecer a lista de verdade
MINIMO_PRODUTOS = 1000


def chave_aliquota(texto_aliquota, sem_impostos, alc):
    """('17,5', None, ' ALC') -> '17.5_alc'; (None, 'SEM IMPOSTOS', None) -> 'sem'."""
    base = "sem" if sem_impostos else texto_aliquota.replace(",", ".")
    return base + ("_alc" if alc else "")


def gtin_valido(codigo):
    """Confere o dígito verificador de um código de barras EAN-13/GTIN-14."""
    if not codigo.isdigit() or len(codigo) not in (8, 12, 13, 14):
        return False
    corpo, dv = codigo[:-1], int(codigo[-1])
    soma = sum(int(d) * (3 if i % 2 == 0 else 1) for i, d in enumerate(reversed(corpo)))
    return (10 - soma % 10) % 10 == dv


@dataclass
class Qualidade:
    """O que foi encontrado e corrigido na lista oficial, com contagens."""
    linhas_notas: int = 0
    produtos: int = 0
    campos_traco: int = 0
    espacos_sobrando: int = 0
    ean_zerado: int = 0
    ean_invalido: int = 0
    ean_repetido: int = 0
    sem_pmc_hospitalar: int = 0
    sem_pmc_outros: int = 0
    precos_asterisco: int = 0
    precos_ilegiveis: int = 0
    exemplos_ilegiveis: list = field(default_factory=list)


def _texto(valor, q):
    if valor is None:
        return ""
    bruto = str(valor)
    limpo = " ".join(bruto.split())
    if RE_VAZIO.match(limpo):
        if limpo:
            q.campos_traco += 1
        return ""
    if limpo != bruto:
        q.espacos_sobrando += 1
    return limpo


def _preco(valor, q, onde):
    """'1234,56' -> 123456 (centavos, inteiro: sem erro de arredondamento)."""
    if valor is None:
        return None, False
    t = str(valor).strip()
    if not t or RE_VAZIO.match(t):
        return None, False
    m = RE_PRECO.match(t)
    if not m:
        q.precos_ilegiveis += 1
        if len(q.exemplos_ilegiveis) < 5:
            q.exemplos_ilegiveis.append(f"{onde}: {t!r}")
        return None, False
    return int(m.group(1)) * 100 + int(m.group(2)), bool(m.group(3))


def _ean(valor, q):
    t = re.sub(r"\D", "", _texto(valor, q))
    if not t:
        return ""
    if set(t) == {"0"}:
        q.ean_zerado += 1
        return ""
    if not gtin_valido(t):
        q.ean_invalido += 1  # mantido: é o que está impresso na caixa
    return t


def ler_lista(caminho):
    """
    Lê o xlsx da CMED. Devolve (info, produtos, qualidade):
    - info: {"publicada_em": "AAAA-MM-DD", "aliquotas": [...]}
    - produtos: lista de dicts, um por apresentação (chave: ggrem)
    """
    wb = load_workbook(caminho, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        linhas = ws.iter_rows(values_only=True)
        q = Qualidade()
        publicada = None
        cabecalho = None
        for i, linha in enumerate(linhas):
            primeira = linha[0] if linha else None
            if isinstance(primeira, str):
                m = RE_PUBLICADA.search(primeira)
                if m and not publicada:
                    publicada = f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
            nomes = [normalizar(c) if c is not None else "" for c in linha]
            if "CODIGO GGREM" in nomes and "PRODUTO" in nomes:
                cabecalho = nomes
                q.linhas_notas = i
                break
            if i > 200:
                break
        if cabecalho is None:
            raise ListaInvalida("cabeçalho (CÓDIGO GGREM, PRODUTO...) não encontrado nas 200 primeiras linhas")

        colunas = {}
        for j, nome in enumerate(cabecalho):
            campo = OBRIGATORIAS.get(nome) or OPCIONAIS.get(nome)
            if campo and campo not in colunas:
                colunas[campo] = j
        faltando = [orig for orig, campo in OBRIGATORIAS.items() if campo not in colunas]
        if faltando:
            raise ListaInvalida(f"colunas obrigatórias ausentes: {', '.join(faltando)}")

        pmc = {}
        for j, nome in enumerate(cabecalho):
            m = RE_PMC.match(nome)
            if m:
                pmc[chave_aliquota(m.group(1), m.group(2), m.group(3))] = j
        if "19" not in pmc or "0" not in pmc:
            raise ListaInvalida(f"colunas de PMC não reconhecidas (achei: {sorted(pmc)})")

        produtos, vistos = [], set()
        for linha in linhas:
            if not linha or all(c is None or str(c).strip() == "" for c in linha):
                continue
            ggrem = re.sub(r"\D", "", str(linha[colunas["ggrem"]] or ""))
            if not ggrem:
                continue  # linha de rodapé/nota no fim, se um dia aparecer
            if ggrem in vistos:
                raise ListaInvalida(f"código GGREM repetido: {ggrem}")
            vistos.add(ggrem)
            p = {"ggrem": ggrem}
            for campo, j in colunas.items():
                if campo == "ggrem":
                    continue
                valor = linha[j] if j < len(linha) else None
                p[campo] = _ean(valor, q) if campo.startswith("ean") else _texto(valor, q)
            p["hospitalar"] = p["hospitalar"].lower() == "sim"
            p["icms0"] = p.get("icms0", "").lower() == "sim"
            p["tarja"] = re.sub(r"^Tarja ", "", p.get("tarja", "")).replace("Sem Tarja", "Sem tarja")
            asterisco = False
            p["pmc"] = {}
            for chave, j in pmc.items():
                centavos, ast = _preco(linha[j] if j < len(linha) else None, q, f"{ggrem} PMC {chave}")
                p["pmc"][chave] = centavos
                asterisco |= ast
            if asterisco:
                q.precos_asterisco += 1
            if p["pmc"]["19"] is None:
                if p["hospitalar"]:
                    q.sem_pmc_hospitalar += 1
                else:
                    q.sem_pmc_outros += 1
            produtos.append(p)
    finally:
        wb.close()

    if len(produtos) < MINIMO_PRODUTOS:
        raise ListaInvalida(f"só {len(produtos)} produtos — a lista oficial tem dezenas de milhares")
    if q.precos_ilegiveis > len(produtos) // 100:
        raise ListaInvalida(f"{q.precos_ilegiveis} preços em formato desconhecido, ex.: {q.exemplos_ilegiveis}")

    contagem = {}
    for p in produtos:
        if p["ean1"]:
            contagem[p["ean1"]] = contagem.get(p["ean1"], 0) + 1
    q.ean_repetido = sum(1 for n in contagem.values() if n > 1)
    q.produtos = len(produtos)
    info = {"publicada_em": publicada, "aliquotas": sorted(pmc, key=_ordem_aliquota)}
    return info, produtos, q


def _ordem_aliquota(chave):
    base = chave.replace("_alc", "")
    return (-1 if base == "sem" else float(base), chave.endswith("_alc"))
