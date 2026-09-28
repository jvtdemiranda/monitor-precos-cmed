"""
Testes do robô (python -m unittest discover -s testes).

A lista de teste imita o arquivo real da CMED: notas antes do cabeçalho,
as mesmas 74 colunas com os mesmos nomes (inclusive os espaços duplos e o
espaço não separável), tudo como texto, preços com vírgula e asterisco,
traços no lugar de campos vazios.
"""
import gzip
import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "robo"))

import coleta  # noqa: E402
import comparar  # noqa: E402
import dados  # noqa: E402
import lista  # noqa: E402
import publicar  # noqa: E402
import rodar  # noqa: E402

ALIQ = ["0", "12", "17", "17,5", "18", "19", "19,5", "20", "20,5", "21", "22", "22,5", "23"]


def cabecalho():
    c = ["SUBSTÂNCIA", "CNPJ", "LABORATÓRIO", "CÓDIGO GGREM", "REGISTRO", "EAN 1", "EAN 2", "EAN 3", "PRODUTO",
         "APRESENTAÇÃO", "CLASSE TERAPÊUTICA", "TIPO DE PRODUTO (STATUS DO PRODUTO)", "REGIME DE PREÇO", "PF Sem Impostos"]
    for tipo in ("PF", "PMC"):
        if tipo == "PMC":
            c.append("PMC Sem Impostos")
        for a in ALIQ:
            c.append(f"{tipo} {a}%" if (tipo, a) == ("PF", "0") else f"{tipo} {a} %")
            if a != "0":
                c.append(f"{tipo} {a} %  ALC")
    return c + ["RESTRIÇÃO HOSPITALAR", "CAP", "CONFAZ 87", "ICMS 0%", "ANÁLISE RECURSAL",
                "LISTA DE CONCESSÃO DE CRÉDITO TRIBUTÁRIO (PIS/COFINS)", "COMERCIALIZAÇÃO 2025", "TARJA", "DESTINAÇÃO COMERCIAL\xa0"]


CAB = cabecalho()


def produto(ggrem, nome, pmc19, ean="7891106000956", hospitalar=False, icms0=False, asterisco=False, **extra):
    """Uma linha no formato do arquivo real. pmc19 em reais (float) ou None."""
    linha = dict.fromkeys(CAB)
    linha.update({"SUBSTÂNCIA": extra.get("substancia", nome), "CNPJ": "18.459.628/0001-15", "LABORATÓRIO": extra.get("lab", "LAB TESTE S.A."),
                  "CÓDIGO GGREM": ggrem, "REGISTRO": "1705600230032", "EAN 1": ean, "EAN 2": "    -     ", "EAN 3": "    -     ",
                  "PRODUTO": nome, "APRESENTAÇÃO": extra.get("apresentacao", "10 MG COM CT BL X 30"), "CLASSE TERAPÊUTICA": "X",
                  "TIPO DE PRODUTO (STATUS DO PRODUTO)": "Genérico", "REGIME DE PREÇO": "Regulado",
                  "RESTRIÇÃO HOSPITALAR": "Sim" if hospitalar else "Não", "ICMS 0%": "Sim" if icms0 else "Não",
                  "TARJA": extra.get("tarja", "Tarja Vermelha")})
    for col in CAB:
        if col.startswith("PMC") and pmc19 is not None:
            fator = 0.9 if "ALC" in col or "Sem" in col else 1.0
            linha[col] = f"{pmc19 * fator:.2f}".replace(".", ",") + ("*" if asterisco else "")
    return [linha[c] for c in CAB]


def lista_xlsx(caminho, linhas, notas=41, publicada="09/09/2026", cab=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    ws.append(["Secretaria Executiva - CMED"])
    ws.append(["LISTA DE PREÇOS DE MEDICAMENTOS - PREÇOS FÁBRICA E MÁXIMOS AO CONSUMIDOR"])
    ws.append([f"Publicada em {publicada} 19h30min."])
    for i in range(notas - 3):
        ws.append([f"nota {i}" if i % 3 else None])
    ws.append(cab or CAB)
    for l in linhas:
        ws.append(l)
    wb.save(caminho)
    return caminho


def basica():
    return [
        produto("538912020009303", "BAYCUTEN N", 50.90, ean="7891106000956", apresentacao="  10 MG/G   CREM X 40 G "),
        produto("505107701157215", "ORENCIA", None, ean="7896016806469", hospitalar=True),
        produto("507619060021902", "VERZENIOS", 5225.85, ean="0000000000000"),
        produto("514518120035014", "HIPOGLOS", 30.53, ean="7891010249908", icms0=True, asterisco=True, tarja="- (*) "),
        produto("514518120035214", "HIPOGLOS", 72.33, ean="7891010249908"),  # mesmo EAN, outro produto
        produto("506418100035404", "BABYMED", 19.18, ean="7896523206944"),  # dígito verificador errado
    ]


class TesteLista(unittest.TestCase):
    def setUp(self):
        self._min = lista.MINIMO_PRODUTOS
        lista.MINIMO_PRODUTOS = 3
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        lista.MINIMO_PRODUTOS = self._min
        self.tmp.cleanup()

    def test_le_e_limpa(self):
        info, ps, q = lista.ler_lista(lista_xlsx(self.dir / "l.xlsx", basica()))
        self.assertEqual(info["publicada_em"], "2026-09-09")
        self.assertIn("19", info["aliquotas"])
        self.assertIn("17.5_alc", info["aliquotas"])
        self.assertEqual(q.linhas_notas, 41)
        p = {x["ggrem"]: x for x in ps}
        self.assertEqual(p["538912020009303"]["pmc"]["19"], 5090)  # centavos
        self.assertEqual(p["538912020009303"]["apresentacao"], "10 MG/G CREM X 40 G")
        self.assertEqual(p["538912020009303"]["ean2"], "")  # "    -     " vira vazio
        self.assertIsNone(p["505107701157215"]["pmc"]["19"])
        self.assertTrue(p["505107701157215"]["hospitalar"])
        self.assertEqual(p["507619060021902"]["ean1"], "")  # "0000000000000" descartado
        self.assertEqual(p["514518120035014"]["pmc"]["19"], 3053)  # asterisco não atrapalha
        self.assertTrue(p["514518120035014"]["icms0"])
        self.assertEqual(p["514518120035014"]["tarja"], "")
        self.assertEqual(p["538912020009303"]["tarja"], "Vermelha")
        self.assertEqual((q.ean_zerado, q.ean_invalido, q.ean_repetido, q.sem_pmc_hospitalar, q.precos_asterisco), (1, 1, 1, 1, 1))
        self.assertGreater(q.espacos_sobrando, 0)

    def test_cabecalho_em_outra_linha(self):
        info, ps, q = lista.ler_lista(lista_xlsx(self.dir / "l.xlsx", basica(), notas=55))
        self.assertEqual(q.linhas_notas, 55)
        self.assertEqual(len(ps), 6)

    def test_coluna_obrigatoria_faltando(self):
        cab = [c if c != "CÓDIGO GGREM" else "CODIGO" for c in CAB]
        with self.assertRaises(lista.ListaInvalida):
            lista.ler_lista(lista_xlsx(self.dir / "l.xlsx", basica(), cab=cab))

    def test_preco_ilegivel_para_tudo(self):
        linhas = basica()
        j = CAB.index("PMC 19 %")
        for l in linhas:
            l[j] = "R$ 12.34"  # formato que a lista nunca usou
        with self.assertRaises(lista.ListaInvalida):
            lista.ler_lista(lista_xlsx(self.dir / "l.xlsx", linhas))

    def test_ggrem_repetido(self):
        # linha idêntica: some; mesmo código com preço diferente: fica o menor PMC
        linhas = basica() + [basica()[0], produto("507619060021902", "VERZENIOS", 4000.00, ean="0000000000000")]
        _, ps, q = lista.ler_lista(lista_xlsx(self.dir / "l.xlsx", linhas))
        self.assertEqual(len(ps), 6)
        self.assertEqual((q.linhas_repetidas, q.ggrem_conflitante), (1, 1))
        self.assertEqual({p["ggrem"]: p for p in ps}["507619060021902"]["pmc"]["19"], 400000)
        self.assertIn("507619060021902", q.exemplos_conflito[0])

    def test_muitos_ggrem_conflitantes_param_tudo(self):
        extras = [produto(f"9{i:014d}", f"P{i}", 10.0) for i in range(6)]
        conflitos = [produto(f"9{i:014d}", f"P{i}", 20.0) for i in range(6)]
        with self.assertRaises(lista.ListaInvalida):  # 6 conflitos: mais que o tolerado (5 ou 1%)
            lista.ler_lista(lista_xlsx(self.dir / "l.xlsx", basica() + extras + conflitos))

    def test_lista_pequena_demais(self):
        lista.MINIMO_PRODUTOS = 1000
        with self.assertRaises(lista.ListaInvalida):
            lista.ler_lista(lista_xlsx(self.dir / "l.xlsx", basica()))

    def test_gtin(self):
        self.assertTrue(lista.gtin_valido("7891106000956"))
        self.assertFalse(lista.gtin_valido("7891106000957"))
        self.assertTrue(lista.gtin_valido("04006381333931"))  # GTIN-14 (EAN-13 válido com zero à esquerda)
        self.assertFalse(lista.gtin_valido("10789800841009"))  # GTIN-14 da lista oficial, com dígito errado


class TesteColeta(unittest.TestCase):
    def setUp(self):
        self.html = (Path(__file__).parent / "fixtures" / "pagina_precos.html").read_text(encoding="utf-8")

    def test_acha_a_lista_de_pmc_na_pagina_real(self):
        data, url = coleta.lista_mais_recente(self.html)
        self.assertEqual(data, "2026-09-09")
        self.assertEqual(url, "https://www.gov.br/anvisa/pt-br/assuntos/medicamentos/cmed/precos/arquivos/"
                              "xls_conformidade_site_20260909_222937320.xlsx/@@download/file")

    def test_ignora_link_de_outro_site(self):
        falso = '<a href="https://exemplo-malicioso.com/xls_conformidade_site_20991231_1.xlsx">PMC - xls</a>'
        data, _ = coleta.lista_mais_recente(self.html + falso)
        self.assertEqual(data, "2026-09-09")

    def test_pagina_mudou(self):
        with self.assertRaises(coleta.LinkNaoEncontrado):
            coleta.lista_mais_recente("<html><body><a href='/outra-coisa.pdf'>PMC</a></body></html>")

    def test_recusa_pagina_de_erro_no_lugar_da_planilha(self):
        with self.assertRaises(ValueError):
            coleta.conferir_planilha(b"<!DOCTYPE html><html>Servico indisponivel</html>")
        self.assertEqual(len(coleta.conferir_planilha(b"PK\x03\x04resto")), 64)


class _Resposta(io.BytesIO):
    """Resposta falsa de urlopen: conteúdo e endereço final (depois de redirecionar)."""
    def __init__(self, corpo, final):
        super().__init__(corpo)
        self.final = final

    def geturl(self):
        return self.final


class TesteDownload(unittest.TestCase):
    URL = "https://www.gov.br/anvisa/pt-br/assuntos/medicamentos/cmed/precos/arquivos/x.xlsx/@@download/file"

    def setUp(self):
        self._orig, self._pausa = coleta.urllib.request.urlopen, coleta.PAUSA
        coleta.PAUSA = 0

    def tearDown(self):
        coleta.urllib.request.urlopen, coleta.PAUSA = self._orig, self._pausa

    def falso(self, corpo=b"PK\x03\x04ok", final=None):
        coleta.urllib.request.urlopen = lambda req, timeout: _Resposta(corpo, final or req.full_url)

    def test_baixa(self):
        self.falso()
        self.assertEqual(coleta.baixar(self.URL), b"PK\x03\x04ok")

    def test_recusa_endereco_fora_da_anvisa(self):
        self.falso()
        for url in ("https://outro-site.com/x.xlsx", "http://www.gov.br/anvisa/x.xlsx"):
            with self.assertRaises(coleta.OrigemRecusada):
                coleta.baixar(url)

    def test_recusa_redirecionamento_para_outro_site(self):
        self.falso(final="https://outro-site.com/x.xlsx")
        with self.assertRaises(coleta.OrigemRecusada):
            coleta.baixar(self.URL)

    def test_recusa_arquivo_gigante(self):
        maximo = coleta.TAMANHO_MAXIMO
        coleta.TAMANHO_MAXIMO = 1000
        try:
            self.falso(corpo=b"PK\x03\x04" + b"0" * 5000)
            with self.assertRaises(coleta.OrigemRecusada):
                coleta.baixar(self.URL)
        finally:
            coleta.TAMANHO_MAXIMO = maximo

    def test_tenta_de_novo_quando_a_conexao_cai(self):
        chamadas = []

        def instavel(req, timeout):
            chamadas.append(1)
            if len(chamadas) < 3:
                raise ConnectionResetError("conexão caiu")
            return _Resposta(b"PK\x03\x04ok", req.full_url)
        coleta.urllib.request.urlopen = instavel
        self.assertEqual(coleta.baixar(self.URL, espera=(0, 0)), b"PK\x03\x04ok")
        self.assertEqual(len(chamadas), 3)


def _p(ggrem, pmc19, nome="X"):
    return {"ggrem": ggrem, "produto": nome, "apresentacao": "", "laboratorio": "", "pmc": {"19": pmc19}}


class TesteComparar(unittest.TestCase):
    def test_mudancas(self):
        ant = [_p("1", 1000), _p("2", 1000), _p("3", 103710), _p("4", None), _p("5", 500)]
        atu = [_p("1", 1000), _p("2", 900), _p("3", 1037109), _p("4", 2000), _p("6", 700)]
        r = comparar.comparar(ant, atu)
        tipos = {m["produto"]["ggrem"]: m for m in r["mudancas"]}
        self.assertNotIn("1", tipos)
        self.assertEqual(tipos["2"]["tipo"], "desceu")
        self.assertFalse(tipos["2"]["fora_do_normal"])
        self.assertTrue(tipos["3"]["fora_do_normal"])  # o caso real do preço 10x maior
        self.assertEqual(r["mudancas"][0]["produto"]["ggrem"], "3")  # fora do normal vem primeiro
        self.assertEqual(tipos["4"]["tipo"], "passou a ter PMC")
        self.assertEqual([p["ggrem"] for p in r["novos"]], ["6"])
        self.assertEqual([p["ggrem"] for p in r["sairam"]], ["5"])

    def test_historico(self):
        precos = {"a": {"1": (0, 100), "2": (0, 100)}, "b": {"1": (0, 105), "2": (0, 100), "3": (0, 5)},
                  "c": {"1": (0, 400), "3": (0, 5)}}
        h = comparar.resumo_historico(["a", "b", "c"], precos.get)
        self.assertNotIn("subiu", h[0])
        self.assertEqual((h[1]["subiu"], h[1]["novos"], h[1]["sairam"]), (1, 1, 0))
        self.assertEqual((h[2]["subiu"], h[2]["fora_do_normal"], h[2]["sairam"]), (1, 1, 1))


class TesteDadosEPublicacao(unittest.TestCase):
    def setUp(self):
        self._min = lista.MINIMO_PRODUTOS
        lista.MINIMO_PRODUTOS = 3
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        lista.MINIMO_PRODUTOS = self._min
        self.tmp.cleanup()

    def processar(self, linhas, publicada):
        arq = lista_xlsx(self.dir / f"{publicada[-4:]}{publicada[3:5]}.xlsx", linhas, publicada=publicada)
        return rodar.processar(arq.read_bytes(), None, "teste", raiz=self.dir / "dados")

    def test_ida_e_volta_e_gz_identico(self):
        self.processar(basica(), "11/08/2026")
        gz = (self.dir / "dados" / "listas" / "2026-08-11.csv.gz").read_bytes()
        self.processar(basica(), "11/08/2026")
        self.assertEqual(gz, (self.dir / "dados" / "listas" / "2026-08-11.csv.gz").read_bytes())
        _, ps = dados.ler_lista_limpa("2026-08-11", self.dir / "dados")
        self.assertEqual(ps[0]["pmc"]["19"], 5090)
        self.assertTrue(ps[1]["hospitalar"])
        self.assertEqual(gzip.decompress(gz)[:5], b"ggrem")

    def test_so_guarda_as_duas_ultimas_listas(self):
        for data in ("21/07/2026", "11/08/2026", "09/09/2026"):
            self.processar(basica(), data)
        self.assertEqual(dados.listas_completas(self.dir / "dados"), ["2026-08-11", "2026-09-09"])
        self.assertEqual(len(dados.publicacoes(self.dir / "dados")), 3)

    def test_recusa_lista_muito_menor_que_a_anterior(self):
        self.processar(basica() * 1 + [produto(f"9{i:014d}", f"P{i}", 10.0) for i in range(20)], "11/08/2026")
        with self.assertRaises(lista.ListaInvalida):
            self.processar(basica(), "09/09/2026")

    def test_publicacao_deterministica_e_sem_formula(self):
        linhas = basica()
        linhas[0][CAB.index("PRODUTO")] = '=HYPERLINK("http://site-falso";"Clique")'
        self.processar(basica(), "11/08/2026")
        self.processar(linhas, "09/09/2026")
        pub = self.dir / "public"
        rodar.publicar_tudo(self.dir / "dados", pub)
        primeira = {p.relative_to(pub): p.read_bytes() for p in pub.rglob("*") if p.is_file()}
        rodar.publicar_tudo(self.dir / "dados", pub)
        segunda = {p.relative_to(pub): p.read_bytes() for p in pub.rglob("*") if p.is_file()}
        self.assertEqual(primeira, segunda)
        wb = load_workbook(pub / "lista-cmed.xlsx")
        for aba in ("Consulta PMC", "PMC por alíquota"):
            celula = next(c for c in wb[aba]["B"] if isinstance(c.value, str) and "HYPERLINK" in c.value)
            self.assertEqual(celula.data_type, "s")
        with zipfile.ZipFile(pub / "lista-cmed.xlsx") as z:
            self.assertNotIn(b"<f>", b"".join(z.read(n) for n in z.namelist() if n.startswith("xl/worksheets/")))

    def test_planilha_arrumada(self):
        self.processar(basica(), "11/08/2026")
        self.processar(basica(), "09/09/2026")
        pub = self.dir / "public"
        rodar.publicar_tudo(self.dir / "dados", pub)
        wb = load_workbook(pub / "lista-cmed.xlsx")
        self.assertEqual(wb.sheetnames, ["Resumo", "Mudanças de preço", "Novos na lista", "Saíram da lista",
                                         "Mês a mês", "Consulta PMC", "PMC por alíquota"])
        for ws in wb.worksheets:
            self.assertFalse(ws.sheet_view.showGridLines, ws.title)
            if ws.title == "Resumo":
                continue
            # cabeçalho fixo e com filtro, logo abaixo do título
            self.assertTrue(ws.freeze_panes, ws.title)
            self.assertTrue(ws.auto_filter.ref, ws.title)
            linha_cab = int(ws.auto_filter.ref.split(":")[0][1:])
            self.assertEqual(int(ws.freeze_panes[1:]), linha_cab + 1, ws.title)
        consulta = wb["Consulta PMC"]
        cab = [c.value for c in consulta[4]][1:]
        self.assertEqual(cab[:2], ["Produto", "Apresentação"])
        self.assertIn("PMC 19% (Pará)", cab)
        self.assertLessEqual(len(cab), 10)  # consulta do dia a dia: poucas colunas
        # preço como número, formatado em reais (não texto)
        preco = next(c for c in consulta["H"][4:] if c.value is not None)
        self.assertIsInstance(preco.value, float)
        self.assertIn("R$", preco.number_format)


if __name__ == "__main__":
    unittest.main()
